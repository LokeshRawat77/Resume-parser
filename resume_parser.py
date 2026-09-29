import os
import time
import json
from pathlib import Path
from pydantic import BaseModel, Field
from dotenv import load_dotenv
from groq import Groq
from pypdf import PdfReader
from docx import Document as DocxDocument

load_dotenv(dotenv_path=os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '.env'))
my_api_key = os.getenv("GROQ_API_KEY")


if not my_api_key:
    raise ValueError("Please set the GROQ_API_KEY environment variable")    

client=Groq(api_key=my_api_key)
model = "openai/gpt-oss-20b"


#Job Description
Job_description = """
Do you enjoy working with data to solve real-world business problems? Are you curious about discovering patterns, trends, and insights that can help organizations make better decisions? Do you enjoy working with large datasets and transforming complex information into meaningful insights?

At our organization, we use data to understand customers, improve products, optimize business processes, and drive strategic decisions. As a Data Analyst, you will work with cross-functional teams to collect, clean, analyze, and visualize data while communicating your findings to business and technical stakeholders.

Our Data Analysts use modern analytical tools and technologies to solve complex business problems. You will work with structured and unstructured data, identify meaningful trends, build dashboards and reports, and provide actionable recommendations based on your analysis.

We are looking for analytical and curious individuals who enjoy asking questions, exploring data, and finding solutions. You will have opportunities to work on challenging datasets, collaborate with experienced professionals, and develop your skills in data analysis, visualization, statistics, and business intelligence.

Key job responsibilities

• Collect, clean, transform, and analyze data from multiple sources to identify meaningful trends and patterns.

• Write SQL queries to extract, manipulate, and analyze data from relational databases.

• Use Python or other programming languages for data cleaning, analysis, and automation.

• Build and maintain dashboards and reports using data visualization tools such as Power BI or Tableau.

• Perform exploratory data analysis to identify trends, anomalies, relationships, and opportunities.

• Work closely with business, product, engineering, and other cross-functional teams to understand analytical requirements.

• Translate business problems into analytical questions and provide data-driven solutions.

• Develop and maintain data reports, dashboards, and analytical models.

• Validate data quality and ensure accuracy and consistency of analytical results.

• Communicate analytical findings clearly through reports, presentations, and visualizations.

• Perform statistical analysis and use appropriate statistical techniques to support business decisions.

• Automate repetitive data analysis and reporting tasks where possible.

• Stay current with emerging technologies, analytics tools, and best practices in the data analytics field.

Basic Qualifications

Bachelor's degree or currently pursuing a bachelor's degree in Data Science, Statistics, Mathematics, Computer Science, Information Systems, Economics, Business Analytics, or a related field.
Experience with SQL and relational databases.
Basic programming experience with Python or another programming language.
Understanding of data structures, data cleaning, and data transformation techniques.
Understanding of basic statistics and analytical concepts.
Experience working with spreadsheets such as Microsoft Excel or Google Sheets.
Basic understanding of data visualization and reporting.
Strong analytical and problem-solving skills.

Preferred Qualifications

Experience with data visualization tools such as Power BI, Tableau, or Looker.
Experience with Python libraries such as Pandas, NumPy, Matplotlib, or Seaborn.
Experience working with large datasets and performing exploratory data analysis.
Knowledge of statistical concepts such as probability, hypothesis testing, correlation, regression, and descriptive statistics.
Experience with cloud platforms such as AWS, Microsoft Azure, or Google Cloud.
Familiarity with data warehouses and databases such as MySQL, PostgreSQL, SQL Server, Snowflake, or BigQuery.
Experience with ETL/ELT processes and data pipelines.
Experience with Git and version control systems.
Experience working on data analytics, business intelligence, or data science projects.
Demonstrated ability to learn new technologies and analytical tools quickly.
Strong written and verbal communication skills.
Ability to explain complex analytical findings to both technical and non-technical stakeholders.
"""

class JobD(BaseModel):
    role:str
    required_skills:list[str]
    preferred_skills:list[str]
    minimun_experience: float| None
    education_requirements:list[str]
    responsibilities: list[str]

jobd_schema = JobD.model_json_schema()

system_prompt = f"""
You are an expart HR assistant .

Your job is to analyze job descriptions and extract 
structured infromation from them.

Return ONLY vaild JSON matching this schema:
{jobd_schema}
IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""

user_propmt=f"""
Analyze the following job description and extract the information according to the schema:
{Job_description}
"""
message_system={
    "role":"system",
    "content":system_prompt
}
message_user={
    "role":"user",
    "content":user_propmt
}
response_fromat={
    "type":"json_object"
}

messages=[message_system,message_user]
response=client.chat.completions.create(model=model,messages=messages,response_format=response_fromat)

answer = response.choices[0].message.content
raw_json = answer
#print(raw_json)
job_data = json.loads(raw_json)
job=JobD(**job_data)

print(job.minimun_experience)
print(job.education_requirements)

#parse real 
class MatchResult(BaseModel):
    score:float
    details:dict
class Experience(BaseModel):
    compay:str |None=None
    role:str |None=None
    duration:str |None=None
    description:str |None=None
    skill_used:list[str]=[]

class Resume(BaseModel):
    name:str| None=None
    email:str| None=None
    phone:str |None=None
    total_experience_years: float | None = None
    skills:list[str]=[]
    experiences:list[Experience]=[]
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []

resume_schema = Resume.model_json_schema()

def parse_resume(resume_text: str):
    resume_system_prompt = f"""
    You are a resume parser. Extract information from the resume strictly based on this schema.
    {resume_schema}
    Return ONLY valid JSON. Do not return the schema itself.
    """
    messages = [
        {"role": "system", "content": resume_system_prompt},
        {"role": "user", "content": f"Parse this resume:\n\n{resume_text}"}
    ]
    response = client.chat.completions.create(
        model=model,
        messages=messages,
        response_format={"type": "json_object"},
        temperature=0
    )
    data = json.loads(response.choices[0].message.content)
    return Resume(**data)

def final_score(job, resume):
    match_schema =MatchResult.model_json_schema()
    prompt =f"""
You are an HR Recruiter.

Compare the candidate's resume with the job description.
JOB DESCRIPITON:
{job.model_dump_json(indent=2)}

CANDIDATE RESUME:
{resume.model_dump_json(indent=2)}
Return JSON matching this schema:

{match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    
"""

    message_system={
        "role" : "system",
        "content" : "You are an HR Recruiter. Return ONLY valid JSON matching the given schema."
    }
    message_user={
        "role" : "user",
        "content" : prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = MatchResult(**data)
    return resume

class Document:
    """Concrete document reader for the resume formats supported here."""

    def __init__(self, file_path):
        self.file_path = file_path
        self.text = self._read()

    def _read(self):
        extension = os.path.splitext(self.file_path)[1].lower()
        if extension == ".pdf":
            return read_pdf(self.file_path)
        if extension == ".docx":
            return read_docx(self.file_path)
        raise ValueError("Unsupported document format; expected .pdf or .docx")

    def extract_text(self):
        """Return the extracted document text."""
        return self.text

    def __str__(self):
        return self.text

    def __len__(self):
        return len(self.text)

def read_pdf(file__path):
    reader = PdfReader(file__path)
    text = ""
    for page in reader.pages:
        text += (page.extract_text() or "") + "\n"
    return text

def read_docx(file_path):
    doc = DocxDocument(file_path)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

    for table in document.tables:
        for row in table.rows:
            if cell.text.strip():
                text += cell.text +"\n"
    return text

def read_resume(file_path):
    if file_path.suffix.lower() ==".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None

#Lets do it now 
resume_folder = Path(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'resumes'))
all_results=[]
for file_path in resume_folder.iterdir():
    # D:\LOKESH_AI\resume_parser\resumes
    if file_path.suffix.lower() not in [".pdf",".docx"]:
        continue
    print("\nProcessing:",file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume = parse_resume(resume_text)  # llm call1
    time.sleep(30)
    result= final_score(job,parsed_resume) #llm call2
    time.sleep(30)
    print("Score:",result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })
    all_results.sort(
        key=lambda candidate: candidate["score"], 
        reverse=True
        )
    top_2 =all_results[:2]
    worst_2 = all_results[-2:]

    print("TOP 2 CANDIDATE")
    for candidate in top_2:
        print(
            f"{candidate['name']}: {candidate['score']}%")
        print(candidate['details'])

    print("WORST 2 CANDIDATE")
    for candidate in worst_2:
        print(
            f"{candidate['name']}: {candidate['score']}%")
        print(candidate['details'])
